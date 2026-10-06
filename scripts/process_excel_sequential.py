"""Sequential Excel backtesting/import processor for NeuroLytics.

Processes one Excel row at a time:
1. Stage 1: persist Open + Jodi-first.
2. Predict Stage 2 using only data completed before the current row.
3. Stage 2: persist the actual Jodi-second + Close.
4. Evaluate the prediction against the actual.
5. Retrain only after an incorrect prediction.
6. Checkpoint after every completed row for safe resume.

The processor deliberately does NOT bulk-import the workbook.
"""
from __future__ import annotations

import argparse
import json
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import Path
from typing import Any
import sys

ROOT_DIR = Path(__file__).resolve().parent.parent
if str(ROOT_DIR) not in sys.path:
    sys.path.insert(0, str(ROOT_DIR))

from openpyxl import load_workbook
from sqlalchemy import func, select

from analytics.live_activity import live_activity
from analytics.sequential_prediction import (
    MIN_TRAIN_ROWS,
    load_status,
    predict_jodi_candidates,
    predict_stage_2,
    train_sequential_models,
)
from analytics.system_orchestrator import NeuroLyticsSystem
from database.engine import SessionLocal
from database.models import HistoricalResult, Market, PredictionFeedback, SequentialPredictionStage

DEFAULT_EXCEL = Path(__file__).resolve().parent.parent / "data" / "RawData.xlsx"
DEFAULT_SHEET = "Open Jodi Close"
DEFAULT_MARKET = "Excel Sequential Market"
DEFAULT_CHECKPOINT = Path(__file__).resolve().parent.parent / "data" / "sequential_excel_checkpoint.json"


@dataclass(frozen=True)
class ExcelRow:
    row_number: int
    result_date: date
    open_result: str
    jodi_result: str
    close_result: str


def _digit_text(value: Any, width: int) -> str:
    if value is None:
        raise ValueError("missing value")
    if isinstance(value, bool):
        raise ValueError("boolean is not a valid digit value")
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    if not text.isdigit() or len(text) > width:
        raise ValueError(f"expected at most {width} digits, got {text!r}")
    return text.zfill(width)


def _jodi_text(value: Any) -> str:
    if value is None:
        raise ValueError("missing Jodi value")
    text = str(value).strip()
    if text.endswith(".0"):
        text = text[:-2]
    if not text.isdigit() or len(text) > 2:
        raise ValueError(f"expected at most 2 Jodi digits, got {text!r}")
    return text.zfill(2)


def _date_value(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return datetime.fromisoformat(str(value).strip().replace("Z", "+00:00")).date()


def read_excel_rows(path: Path, sheet_name: str = DEFAULT_SHEET) -> list[ExcelRow]:
    if not path.exists():
        raise FileNotFoundError(f"Excel file not found: {path}")
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        if sheet_name not in workbook.sheetnames:
            raise ValueError(f"Worksheet not found: {sheet_name!r}")
        sheet = workbook[sheet_name]
        rows: list[ExcelRow] = []
        seen_dates: set[date] = set()
        previous_date: date | None = None
        for row_number, values in enumerate(sheet.iter_rows(min_row=2, values_only=True), start=2):
            if not any(value is not None for value in values):
                continue
            if len(values) < 4:
                raise ValueError(f"row {row_number}: expected Date/Open/Jodi/Close columns")
            result_date = _date_value(values[0])
            if result_date in seen_dates:
                raise ValueError(f"row {row_number}: duplicate date {result_date}")
            if previous_date is not None and result_date <= previous_date:
                raise ValueError(f"row {row_number}: dates are not strictly increasing")
            seen_dates.add(result_date)
            previous_date = result_date
            rows.append(
                ExcelRow(
                    row_number=row_number,
                    result_date=result_date,
                    open_result=_digit_text(values[1], 3),
                    jodi_result=_jodi_text(values[2]),
                    close_result=_digit_text(values[3], 3),
                )
            )
        if not rows:
            raise ValueError("worksheet contains no data rows")
        return rows
    finally:
        workbook.close()


def _checkpoint_load(path: Path) -> dict[str, Any]:
    if not path.exists():
        return {"version": 1, "source": None, "market": None, "last_row_number": 1, "processed": 0}
    return json.loads(path.read_text(encoding="utf-8"))


def _checkpoint_save(path: Path, *, source: str, market: str, row: ExcelRow, processed: int) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "version": 1,
        "source": source,
        "market": market,
        "last_row_number": row.row_number,
        "last_date": row.result_date.isoformat(),
        "processed": processed,
        "updated_at": datetime.now().astimezone().isoformat(),
    }
    temporary = path.with_suffix(path.suffix + ".tmp")
    temporary.write_text(json.dumps(payload, indent=2, sort_keys=True), encoding="utf-8")
    temporary.replace(path)


def _existing_result(market_name: str, result_date: date) -> bool:
    db = SessionLocal()
    try:
        market = db.scalar(select(Market).where(Market.name == market_name))
        if market is None:
            return False
        return db.scalar(
            select(HistoricalResult.id).where(
                HistoricalResult.market_id == market.id,
                HistoricalResult.result_date == result_date,
            )
        ) is not None
    finally:
        db.close()


def _historical_count(market_name: str) -> int:
    """Return completed historical rows available for the sequential model."""
    db = SessionLocal()
    try:
        market = db.scalar(select(Market).where(Market.name == market_name))
        if market is None:
            return 0
        return int(
            db.scalar(
                select(func.count(HistoricalResult.id)).where(
                    HistoricalResult.market_id == market.id,
                )
            )
            or 0
        )
    finally:
        db.close()


def _top_prediction(prediction: dict[str, Any]) -> tuple[str, str]:
    jodi = prediction["jodi_second_candidates"][0]["digit"]
    close = "".join(
        str(prediction["close_digit_candidates"][position][0]["digit"])
        for position in ("first", "second", "third")
    )
    return str(jodi), close


def _record_jodi_feedback(
    *,
    market_id: int,
    result_date: date,
    predicted_first: str,
    predicted_second: str,
    actual_first: str,
    actual_second: str,
    model_identity: str | None,
) -> tuple[bool, bool]:
    """Persist separate B.1 Jodi digit evaluations.

    Each Jodi digit is evaluated independently. A failed first or second
    digit marks the row as requiring retraining, but retraining itself is
    performed once after the complete row is available.
    """
    results = (
        (
            "JODI_FIRST",
            predicted_first,
            actual_first,
            "jodi_first_top1_exact",
        ),
        (
            "JODI_SECOND",
            predicted_second,
            actual_second,
            "jodi_second_top1_exact",
        ),
    )
    db = SessionLocal()
    try:
        outcomes: list[bool] = []
        for stage, predicted, actual, evaluation in results:
            is_correct = predicted == actual
            outcomes.append(is_correct)
            db.add(
                PredictionFeedback(
                    market_id=market_id,
                    result_date=result_date,
                    stage=stage,
                    predicted_value=str(predicted),
                    actual_value=str(actual),
                    is_correct=is_correct,
                    model_identity=model_identity,
                    notes=f"evaluation={evaluation};independent_b1=True",
                )
            )
        db.commit()
    finally:
        db.close()

    for stage, predicted, actual, _ in results:
        is_correct = predicted == actual
        live_activity.emit(
            "Jodi digit prediction evaluated",
            status="RUNNING",
            current_date=result_date.isoformat(),
            stage=stage,
            operation="Evaluating independent Jodi digit",
            prediction=str(predicted),
            actual=str(actual),
            result="CORRECT" if is_correct else "INCORRECT",
            model_version=model_identity,
        )

    return outcomes[0], outcomes[1]


def _record_feedback(
    *,
    market_id: int,
    result_date: date,
    predicted_jodi: str,
    predicted_close: str,
    actual_jodi: str,
    actual_close: str,
    model_identity: str | None,
) -> bool:
    jodi_correct = predicted_jodi == actual_jodi
    close_correct = predicted_close == actual_close
    is_correct = jodi_correct and close_correct
    predicted = f"{predicted_jodi} {predicted_close}"
    actual = f"{actual_jodi} {actual_close}"

    db = SessionLocal()
    try:
        feedback = PredictionFeedback(
            market_id=market_id,
            result_date=result_date,
            stage="STAGE_2",
            predicted_value=predicted,
            actual_value=actual,
            is_correct=is_correct,
            model_identity=model_identity,
            notes=(
                f"jodi_correct={jodi_correct};"
                f"close_correct={close_correct};"
                "evaluation=top1_exact"
            ),
        )
        db.add(feedback)
        db.commit()
    finally:
        db.close()

    snapshot = live_activity.snapshot()["state"]
    total_predictions = snapshot["total_predictions"] + 1
    correct_predictions = snapshot["correct_predictions"] + (1 if is_correct else 0)
    incorrect_predictions = snapshot["incorrect_predictions"] + (0 if is_correct else 1)
    live_activity.emit(
        "Prediction evaluated",
        status="RUNNING",
        current_date=result_date.isoformat(),
        stage="STAGE_2",
        operation="Comparing prediction with actual",
        prediction=predicted,
        actual=actual,
        result="CORRECT" if is_correct else "INCORRECT",
        total_predictions=total_predictions,
        correct_predictions=correct_predictions,
        incorrect_predictions=incorrect_predictions,
        accuracy=(correct_predictions / total_predictions) * 100.0,
        model_version=model_identity,
    )
    return is_correct


def _market_id(market_name: str) -> int:
    db = SessionLocal()
    try:
        market = db.scalar(select(Market).where(Market.name == market_name))
        if market is None:
            raise ValueError(f"market {market_name!r} does not exist")
        return int(market.id)
    finally:
        db.close()


def process_excel(
    excel_path: Path = DEFAULT_EXCEL,
    *,
    market_name: str = DEFAULT_MARKET,
    checkpoint_path: Path = DEFAULT_CHECKPOINT,
    start_row: int | None = None,
    limit: int | None = None,
    dry_run: bool = False,
) -> dict[str, Any]:
    rows = read_excel_rows(excel_path)
    checkpoint = _checkpoint_load(checkpoint_path)
    source_key = str(excel_path.resolve())

    if checkpoint.get("source") not in (None, source_key):
        raise ValueError("checkpoint belongs to a different Excel source; use a new checkpoint file")
    if checkpoint.get("market") not in (None, market_name):
        raise ValueError("checkpoint belongs to a different market")

    resume_after = int(checkpoint.get("last_row_number", 1))
    first_row = start_row if start_row is not None else resume_after + 1
    selected = [row for row in rows if row.row_number >= first_row]
    if limit is not None:
        selected = selected[:limit]

    live_activity.reset(source="EXCEL_SEQUENTIAL")
    live_activity.emit(
        "Starting sequential Excel processing",
        source="EXCEL_SEQUENTIAL",
        operation="Starting sequential Excel processing",
        processed_entries=int(checkpoint.get("processed", 0)),
    )

    if dry_run:
        live_activity.set_status("READY", "Dry-run validation completed")
        return {
            "status": "VALID",
            "dry_run": True,
            "source": source_key,
            "rows": len(rows),
            "selected_rows": len(selected),
            "first_selected_row": selected[0].row_number if selected else None,
            "last_selected_row": selected[-1].row_number if selected else None,
        }

    system = NeuroLyticsSystem()
    processed = int(checkpoint.get("processed", 0))
    predictions = 0
    correct = 0
    incorrect = 0
    retrained = 0
    warmup = 0

    for row in selected:
        live_activity.emit(
            "Processing Excel entry",
            status="RUNNING",
            source="EXCEL_SEQUENTIAL",
            current_date=row.result_date.isoformat(),
            stage="STAGE_1",
            operation=f"Reading Excel row {row.row_number}",
            actual=f"{row.jodi_result} {row.close_result}",
        )

        if _existing_result(market_name, row.result_date):
            raise ValueError(
                f"row {row.row_number} ({row.result_date}) already exists in database; "
                "refusing to skip it silently"
            )

        jodi_first = row.jodi_result[0]
        jodi_second = row.jodi_result[1]

        # Phase B.1 prediction must happen BEFORE the current row's Jodi
        # result is stored. This preserves the temporal boundary.
        model_status = load_status()
        available_history = _historical_count(market_name)
        jodi_prediction: dict[str, Any] | None = None
        predicted_jodi_first: str | None = None
        predicted_jodi_second: str | None = None

        if model_status.get("status") == "COMPLETED" and available_history >= MIN_TRAIN_ROWS:
            jodi_prediction = predict_jodi_candidates(
                row.open_result,
                target_date=row.result_date,
                top_k=10,
                market_name=market_name,
            )
            top_candidate = jodi_prediction["candidates"][0]
            predicted_jodi_first = str(top_candidate.jodi[0])
            predicted_jodi_second = str(top_candidate.jodi[1])
            live_activity.emit(
                "Jodi B.2 ranking generated",
                current_date=row.result_date.isoformat(),
                stage="JODI_B2",
                operation="Ranked Jodi candidates and enforced family constraints",
                prediction=top_candidate.jodi,
                model_version=jodi_prediction.get("ranking_identity"),
            )
        else:
            warmup += 1
            live_activity.emit(
                "Jodi model warm-up entry",
                current_date=row.result_date.isoformat(),
                stage="JODI_B1",
                operation=(
                    f"Prediction deferred until {MIN_TRAIN_ROWS} "
                    "historical rows are available"
                ),
            )

        stage = system.save_stage1(
            row.result_date,
            market_name,
            row.open_result,
            jodi_first,
        )
        live_activity.emit(
            "Stage 1 stored in database",
            current_date=row.result_date.isoformat(),
            stage="STAGE_1",
            operation="Generating Stage 2 prediction",
            actual=None,
        )

        model_status = load_status()
        prediction: dict[str, Any] | None = None
        if model_status.get("status") == "COMPLETED" and available_history >= MIN_TRAIN_ROWS:
            prediction = predict_stage_2(
                row.open_result,
                jodi_first,
                top_k=10,
                market_id=stage.payload["market_id"],
                market_name=market_name,
            )
            predictions += 1
            predicted_jodi, predicted_close = _top_prediction(prediction)
            live_activity.emit(
                "Stage 2 prediction generated",
                current_date=row.result_date.isoformat(),
                stage="STAGE_2",
                operation="Waiting for actual Jodi-second + Close",
                prediction=f"{predicted_jodi} {predicted_close}",
            )
        else:
            live_activity.emit(
                "Stage 2 model warm-up entry",
                current_date=row.result_date.isoformat(),
                stage="STAGE_2",
                operation=f"Prediction deferred until {MIN_TRAIN_ROWS} historical rows are available",
            )

        completed = system.complete_stage2(stage.payload["stage_id"], jodi_second, row.close_result)
        live_activity.emit(
            "Stage 2 actual result stored",
            current_date=row.result_date.isoformat(),
            stage="STAGE_2",
            operation="Evaluating prediction",
            actual=f"{jodi_second} {row.close_result}",
        )

        row_requires_retraining = False

        if jodi_prediction is not None:
            first_correct, second_correct = _record_jodi_feedback(
                market_id=completed.payload["market_id"],
                result_date=row.result_date,
                predicted_first=predicted_jodi_first or "",
                predicted_second=predicted_jodi_second or "",
                actual_first=jodi_first,
                actual_second=jodi_second,
                model_identity=(
                    jodi_prediction.get("model_identities", {})
                    .get("model_identity")
                ),
            )
            if not first_correct or not second_correct:
                row_requires_retraining = True

        if prediction is not None:
            predicted_jodi, predicted_close = _top_prediction(prediction)
            is_correct = _record_feedback(
                market_id=completed.payload["market_id"],
                result_date=row.result_date,
                predicted_jodi=predicted_jodi,
                predicted_close=predicted_close,
                actual_jodi=jodi_second,
                actual_close=row.close_result,
                model_identity=prediction.get("model_identity"),
            )
            if is_correct:
                correct += 1
            else:
                incorrect += 1
                row_requires_retraining = True

        if jodi_prediction is not None or prediction is not None:
            if row_requires_retraining:
                retrain_result = train_sequential_models()
                retrained += 1
                live_activity.emit(
                    "Sequential model retraining completed",
                    current_date=row.result_date.isoformat(),
                    stage="MODEL",
                    operation="Ready for next Excel entry",
                    retraining_events=retrained,
                )
                if retrain_result.get("status") != "COMPLETED":
                    raise RuntimeError(
                        f"Retraining failed after row {row.row_number}: {retrain_result}"
                    )
        elif model_status.get("status") != "COMPLETED":
            # Bootstrap only after the current row has been fully completed.
            completed_count = int(
                system.database_snapshot()["historical_results"]
            )
            if completed_count >= MIN_TRAIN_ROWS:
                bootstrap = train_sequential_models()
                if bootstrap.get("status") != "COMPLETED":
                    raise RuntimeError(
                        f"Initial sequential training failed after row {row.row_number}: {bootstrap}"
                    )
                live_activity.emit(
                    "Sequential model initial training completed",
                    current_date=row.result_date.isoformat(),
                    stage="MODEL",
                    operation="Model ready for next Excel entry",
                )

        processed += 1
        _checkpoint_save(
            checkpoint_path,
            source=source_key,
            market=market_name,
            row=row,
            processed=processed,
        )
        snapshot = live_activity.snapshot()["state"]
        live_activity.emit(
            "Entry evaluation completed",
            status="READY",
            current_date=row.result_date.isoformat(),
            stage="COMPLETE",
            operation=f"Completed Excel row {row.row_number}; ready for next entry",
            processed_entries=processed,
            total_predictions=snapshot["total_predictions"],
            correct_predictions=snapshot["correct_predictions"],
            incorrect_predictions=snapshot["incorrect_predictions"],
            retraining_events=retrained,
        )

    final_status = "COMPLETED" if not selected else "COMPLETED"
    live_activity.emit(
        "Sequential Excel processing completed",
        status="READY",
        source="EXCEL_SEQUENTIAL",
        operation="Sequential Excel processing completed",
        processed_entries=processed,
    )
    return {
        "status": final_status,
        "source": source_key,
        "market": market_name,
        "rows_in_excel": len(rows),
        "rows_processed_this_run": len(selected),
        "processed_total": processed,
        "predictions": predictions,
        "correct_predictions": correct,
        "incorrect_predictions": incorrect,
        "retraining_events": retrained,
        "warmup_entries": warmup,
        "checkpoint": str(checkpoint_path),
    }


def main() -> int:
    parser = argparse.ArgumentParser(description="Process RawData.xlsx sequentially through NeuroLytics.")
    parser.add_argument("--excel", type=Path, default=DEFAULT_EXCEL)
    parser.add_argument("--market", default=DEFAULT_MARKET)
    parser.add_argument("--checkpoint", type=Path, default=DEFAULT_CHECKPOINT)
    parser.add_argument("--start-row", type=int)
    parser.add_argument("--limit", type=int)
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    result = process_excel(
        args.excel,
        market_name=args.market,
        checkpoint_path=args.checkpoint,
        start_row=args.start_row,
        limit=args.limit,
        dry_run=args.dry_run,
    )
    print(json.dumps(result, indent=2, sort_keys=True))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
